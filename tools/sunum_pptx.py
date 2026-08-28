"""Sunum PDF'inden PPTX üretir — şartname madde 6 «sunum materyali» kalemi.

    make sunum-pptx        # docs/sunum/Svartal_Sunum.pdf -> ...pptx

NEDEN PDF'TEN, HTML'DEN DEĞİL:
    Slaytların tek kaynağı `docs/sunum/sunum.html`. Aynı tasarımı ikinci kez
    PowerPoint kutularıyla kurmak, iki slayt takımı demektir; biri düzeltilir
    öbürü unutulur ve jüri iki farklı sayı görür. Bu yüzden PPTX **türetilmiş**
    bir çıktıdır: her sayfa, PDF'in o sayfasının tam kanama (full-bleed)
    görüntüsüdür. Tek kaynak HTML olarak kalır.

    Bedeli açıktır ve bilerek ödeniyor: PPTX içindeki metin seçilemez. Şartname
    biçim istiyor, düzenlenebilirlik istemiyor. Düzenlenmesi gereken şey
    HTML'dir.

NEDEN KONUŞMA METNİ GÖMÜLÜYOR:
    `KONUSMA_METNI.md` sunum sırasında kimin ne söyleyeceğini taşıyor. PowerPoint
    açan kişi (jüri salonundaki makine bizimki olmayabilir) notları slaytın
    altında görür; ayrı bir dosya taşımak zorunda kalmaz.

RASTERLEŞTİRİCİ SEÇİMİ — LİSANS:
    `pypdfium2` (Apache-2.0 / BSD-3-Clause, PDFium). Yaygın alternatif
    `PyMuPDF` **AGPL-3.0**'dır; Apache-2.0 lisanslı bir teslimatın bağımlılık
    ağacına girmesi lisans uyum tablosunu (docs/LISANSLAR.md) bozardı.
"""

from __future__ import annotations

import hashlib
import io
import re
import sys
from pathlib import Path

KOK = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KOK))

PDF = KOK / "docs" / "sunum" / "Svartal_Sunum.pdf"
NOTLAR = KOK / "docs" / "sunum" / "KONUSMA_METNI.md"
CIKTI = KOK / "docs" / "sunum" / "Svartal_Sunum.pptx"

OLCEK = 2.5
"""Rasterleştirme ölçeği. 960×540 pt sayfa → 2400×1350 piksel.

Projeksiyon 1920×1080'dir; bunun üstü dosyayı büyütmekten başka bir şey
yapmaz, altı ise yazıları jüri ekranında bulanıklaştırır."""


def _notlar() -> dict[int, str]:
    """Sayfa numarası -> konuşma notu. Başlık biçimi: `## 02 · Problem …`."""
    if not NOTLAR.exists():
        return {}
    metin = NOTLAR.read_text(encoding="utf-8")
    bulunan: dict[int, str] = {}
    bolumler = re.split(r"^## ", metin, flags=re.MULTILINE)[1:]
    for bolum in bolumler:
        baslik, _, govde = bolum.partition("\n")
        eslesme = re.match(r"\s*(\d{1,2})\s*·", baslik)
        if not eslesme:
            continue  # «Canlı demo» gibi sayfaya bağlanmayan bölümler
        sayfa = int(eslesme.group(1))
        temiz = baslik.replace("&nbsp;", " ").strip()
        govde = re.sub(r"^>\s?", "", govde, flags=re.MULTILINE).strip()
        bulunan[sayfa] = f"{temiz}\n\n{govde}"
    return bulunan


def _pdf_ozeti() -> str:
    """Kaynak PDF'in sha256'sı — PPTX'in hangi PDF'ten üretildiğinin kanıtı."""
    return hashlib.sha256(PDF.read_bytes()).hexdigest()


def uret() -> Path:
    import pypdfium2 as pdfium
    from pptx import Presentation
    from pptx.util import Emu, Inches

    if not PDF.exists():
        raise SystemExit(f"PDF yok: {PDF} — önce `make sunum` koşun.")

    belge = pdfium.PdfDocument(PDF)
    genislik_pt, yukseklik_pt = belge[0].get_size()

    sunum = Presentation()
    # Slayt ölçüsü PDF'in kendi oranından gelir; 16:9 varsayımı yazılmaz.
    sunum.slide_width = Emu(int(Inches(genislik_pt / 72).emu))
    sunum.slide_height = Emu(int(Inches(yukseklik_pt / 72).emu))
    bos_duzen = sunum.slide_layouts[6]  # «Blank» — yer tutucu kutusu yok

    notlar = _notlar()
    for sira, sayfa in enumerate(belge, start=1):
        resim = sayfa.render(scale=OLCEK).to_pil()
        tampon = io.BytesIO()
        resim.save(tampon, format="PNG")
        tampon.seek(0)

        slayt = sunum.slides.add_slide(bos_duzen)
        slayt.shapes.add_picture(
            tampon, 0, 0, width=sunum.slide_width, height=sunum.slide_height
        )
        if sira in notlar:
            slayt.notes_slide.notes_text_frame.text = notlar[sira]

    # KAYNAK PDF'İN ÖZETİ DOSYAYA YAZILIR — bayatlık nöbetçisinin dayanağı.
    # Dosya tarihine bakmak işe yaramaz: `git checkout` her dosyaya o anın
    # tarihini yazar, yani depodan taze çekilmiş bayat bir PPTX «yeni» görünür.
    # İçerik özeti klonlamadan etkilenmez (`tests/test_sunum_sayilari.py`).
    sunum.core_properties.comments = f"kaynak-pdf-sha256={_pdf_ozeti()}"
    sunum.save(CIKTI)
    return CIKTI


def main() -> int:
    yol = uret()
    boyut = yol.stat().st_size / 1024
    print(f"✅ {yol.relative_to(KOK)} ({boyut:.0f} KB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
